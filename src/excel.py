from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


def convertir_fecha(fecha):
    """Convierte una fecha de la OC a una fecha real."""

    if not fecha:
        return None

    try:
        return datetime.strptime(fecha, "%d-%m-%Y")
    except ValueError:
        return None


def convertir_decimal(valor):
    """Convierte Decimal a un valor numérico compatible con Excel."""

    if valor is None:
        return None

    return float(valor)


def crear_fila_orden(
    datos,
    totales,
    moneda,
    simbolo_moneda,
    archivo_origen,
    numero_paginas,
    estado_proceso,
    error_proceso,
    totales_validos,
    articulos_validos,
):
    """Prepara una OC para la hoja ORDENES."""

    return {
        "numero_oc": datos["numero_oc"],
        "codigo_hospital": datos["codigo_hospital"],
        "hospital": datos["hospital"],
        "fecha": convertir_fecha(datos["fecha"]),
        "proveedor": datos["proveedor"],
        "rut_proveedor": datos["rut_proveedor"],
        "direccion_proveedor": datos["direccion_proveedor"],
        "ciudad": datos["ciudad"],
        "vendedor": datos["vendedor"],
        "telefono": datos["telefono"],
        "email": datos["email"],
        "forma_pago": datos["forma_pago"],
        "solicitud": datos["solicitud"],
        "nota_interna": datos["nota_interna"],
        "solicitante_oc": datos["solicitante_oc"],
        "primer_aprobador": datos["primer_aprobador"],
        "moneda": moneda,
        "simbolo_moneda": simbolo_moneda,
        "subtotal": convertir_decimal(totales["subtotal"]),
        "descuento": convertir_decimal(totales["descuento"]),
        "cargos": convertir_decimal(totales["cargos"]),
        "neto": convertir_decimal(totales["neto"]),
        "iva": convertir_decimal(totales["iva"]),
        "otros_impuestos": convertir_decimal(
            totales["otros_impuestos"]
        ),
        "total": convertir_decimal(totales["total"]),
        "archivo_origen": archivo_origen,
        "estado_proceso": estado_proceso,
        "error_proceso": error_proceso,
        "fecha_procesamiento": datetime.now(),
        "numero_paginas": numero_paginas,
        "totales_validos": totales_validos,
        "articulos_validos": articulos_validos,
    }


def crear_filas_detalle(numero_oc, articulos, moneda):
    """Prepara los artículos para la hoja DETALLE_OC."""

    filas = []

    for articulo in articulos:
        filas.append(
            {
                "numero_oc": numero_oc,
                "numero_item": articulo["numero"],
                "codigo": articulo["codigo"],
                "descripcion": articulo["descripcion"],
                "unidad_medida": articulo["unidad"],
                "cantidad": convertir_decimal(
                    articulo["cantidad"]
                ),
                "precio_unitario": convertir_decimal(
                    articulo["precio_unitario"]
                ),
                "total_linea": convertir_decimal(
                    articulo["total_linea"]
                ),
                "moneda": moneda,
            }
        )

    return filas


def escribir_hoja(hoja, filas):
    """Escribe encabezados y datos en una hoja."""

    if not filas:
        return

    encabezados = list(filas[0].keys())
    hoja.append(encabezados)

    for fila in filas:
        hoja.append(
            [
                fila.get(encabezado, "")
                for encabezado in encabezados
            ]
        )


def aplicar_formato_encabezado(hoja):
    """Aplica formato a los encabezados."""

    relleno = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    for celda in hoja[1]:
        celda.font = Font(
            bold=True,
            color="FFFFFF",
        )
        celda.fill = relleno
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )


def ajustar_columnas(hoja):
    """Ajusta el ancho de las columnas."""

    for columna in hoja.columns:
        letra = columna[0].column_letter

        ancho_maximo = 0

        for celda in columna:
            if celda.value is not None:
                largo = len(str(celda.value))
                ancho_maximo = max(ancho_maximo, largo)

        ancho = min(ancho_maximo + 2, 45)
        hoja.column_dimensions[letra].width = ancho


def aplicar_formatos_ordenes(hoja):
    """Aplica formatos específicos a la hoja ORDENES."""

    encabezados = {
        celda.value: celda.column
        for celda in hoja[1]
    }

    columna_fecha = encabezados.get("fecha")
    columna_procesamiento = encabezados.get(
        "fecha_procesamiento"
    )

    columnas_monetarias = [
        "subtotal",
        "descuento",
        "cargos",
        "neto",
        "iva",
        "otros_impuestos",
        "total",
    ]

    for fila in range(2, hoja.max_row + 1):

        if columna_fecha:
            hoja.cell(
                row=fila,
                column=columna_fecha,
            ).number_format = "dd-mm-yyyy"

        if columna_procesamiento:
            hoja.cell(
                row=fila,
                column=columna_procesamiento,
            ).number_format = "dd-mm-yyyy hh:mm"

        moneda = hoja.cell(
            row=fila,
            column=encabezados["moneda"],
        ).value

        for nombre_columna in columnas_monetarias:
            columna = encabezados.get(nombre_columna)

            if not columna:
                continue

            celda = hoja.cell(
                row=fila,
                column=columna,
            )

            if moneda == "CLP":
                celda.number_format = '#,##0'

            elif moneda in ("USD", "EUR"):
                celda.number_format = '#,##0.00'

            elif moneda == "UF":
                celda.number_format = '#,##0.0000'


def aplicar_formatos_detalle(hoja):
    """Aplica formatos numéricos a DETALLE_OC."""

    encabezados = {
        celda.value: celda.column
        for celda in hoja[1]
    }

    for fila in range(2, hoja.max_row + 1):

        moneda = hoja.cell(
            row=fila,
            column=encabezados["moneda"],
        ).value

        celda_cantidad = hoja.cell(
            row=fila,
            column=encabezados["cantidad"],
        )

        celda_precio = hoja.cell(
            row=fila,
            column=encabezados["precio_unitario"],
        )

        celda_total = hoja.cell(
            row=fila,
            column=encabezados["total_linea"],
        )

        celda_cantidad.number_format = '#,##0.00'

        if moneda == "CLP":
            celda_precio.number_format = '#,##0'
            celda_total.number_format = '#,##0'

        elif moneda in ("USD", "EUR"):
            celda_precio.number_format = '#,##0.00'
            celda_total.number_format = '#,##0.00'

        elif moneda == "UF":
            celda_precio.number_format = '#,##0.0000'
            celda_total.number_format = '#,##0.0000'


def crear_tabla(hoja, nombre):
    """Convierte los datos de la hoja en una tabla de Excel."""

    if hoja.max_row < 2:
        return

    referencia = (
        f"A1:"
        f"{hoja.cell(row=1, column=hoja.max_column).column_letter}"
        f"{hoja.max_row}"
    )

    tabla = Table(
        displayName=nombre,
        ref=referencia,
    )

    estilo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    tabla.tableStyleInfo = estilo
    hoja.add_table(tabla)


def exportar_excel(filas_ordenes, filas_detalle, ruta_salida):
    """Genera el archivo Excel final."""

    ruta_salida = Path(ruta_salida)

    # Crea la carpeta si todavía no existe.
    ruta_salida.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    libro = Workbook()

    hoja_ordenes = libro.active
    hoja_ordenes.title = "ORDENES"

    hoja_detalle = libro.create_sheet(
        "DETALLE_OC"
    )

    escribir_hoja(
        hoja_ordenes,
        filas_ordenes,
    )

    escribir_hoja(
        hoja_detalle,
        filas_detalle,
    )

    aplicar_formato_encabezado(
        hoja_ordenes
    )

    aplicar_formato_encabezado(
        hoja_detalle
    )

    aplicar_formatos_ordenes(
        hoja_ordenes
    )

    aplicar_formatos_detalle(
        hoja_detalle
    )

    ajustar_columnas(
        hoja_ordenes
    )

    ajustar_columnas(
        hoja_detalle
    )

    # Mantiene visibles los encabezados al desplazarse.
    hoja_ordenes.freeze_panes = "A2"
    hoja_detalle.freeze_panes = "A2"

    # Permite visualizar mejor textos extensos.
    for hoja in (
        hoja_ordenes,
        hoja_detalle,
    ):
        for fila in hoja.iter_rows():
            for celda in fila:
                celda.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    crear_tabla(
        hoja_ordenes,
        "TablaOrdenes",
    )

    crear_tabla(
        hoja_detalle,
        "TablaDetalleOC",
    )

    libro.save(ruta_salida)